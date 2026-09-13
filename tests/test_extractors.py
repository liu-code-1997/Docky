import pytest
from rag.extractors import extract_text, supported_suffixes


def test_txt_and_md_extract_plain(tmp_path):
    (tmp_path / "a.txt").write_text("hello txt", encoding="utf-8")
    assert extract_text(tmp_path / "a.txt") == "hello txt"
    (tmp_path / "b.md").write_text("# T\nbody here", encoding="utf-8")
    assert "body here" in extract_text(tmp_path / "b.md")


def test_html_strips_tags_and_scripts(tmp_path):
    (tmp_path / "a.html").write_text(
        "<html><head><style>x{color:red}</style>"
        "<script>var y=1</script></head>"
        "<body><h1>Title</h1><p>Real content</p></body></html>",
        encoding="utf-8")
    out = extract_text(tmp_path / "a.html")
    assert "Real content" in out and "Title" in out
    assert "var y" not in out and "color:red" not in out


def test_docx_extracts_paragraphs(tmp_path):
    from docx import Document
    d = Document(); d.add_paragraph("First para"); d.add_paragraph("Second para")
    d.save(str(tmp_path / "a.docx"))
    out = extract_text(tmp_path / "a.docx")
    assert "First para" in out and "Second para" in out


def test_pdf_is_supported():
    assert ".pdf" in supported_suffixes()


def test_unsupported_suffix_raises(tmp_path):
    (tmp_path / "a.xyz").write_text("x", encoding="utf-8")
    with pytest.raises(ValueError):
        extract_text(tmp_path / "a.xyz")
